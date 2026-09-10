# -*- coding: utf-8 -*-
"""Wykaz zestawów BEZ opisu redakcyjnego, od najnowszych po najstarsze.

Uzupełnienie kolejki redakcyjnej (scripts/kolejka-redakcyjna.py), która sortuje
po cenie katalogowej. Ten arkusz odpowiada na inne pytanie: co z rzeczy nowych
nie ma jeszcze żadnego tekstu — bo świeży zestaw bez opisu traci ruch wtedy,
kiedy ludzie go szukają, a nie dwa lata później.

Definicja "bez opisu redakcyjnego" jest ta sama co status 'do opisania'
w kolejce: ani karty Piotra, ani naszego opisu, ani person.
"""
import json, datetime
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BAZA = '/home/user/blogoklockach/'
kat = json.load(open(BAZA + 'src/data/katalog.json'))
sety = json.load(open(BAZA + 'src/data/sety.json'))
karty = {k for k in json.load(open(BAZA + 'src/data/karty_setow.json')) if k != '_meta'}
feed = json.load(open(BAZA + 'src/data/oferty_feed.json')).get('sety', {})

def bez_opisu(nr):
    w = sety.get(nr)
    return not (nr in karty or (w and (w.get('opis') or (w.get('dla_rodzica') and w.get('dla_afol')))))

rows = [(s, r) for s, v in kat.items() if s != '_meta' and isinstance(v, list) for r in v]
kand = [(s, r) for s, r in rows
        if isinstance(r.get('rok'), int) and 2020 <= r['rok'] <= 2026
        and r.get('status') == 'dostepny'
        and r.get('cena_katalogowa') not in (None, '', 0)
        and bez_opisu(str(r['numer']))]

best = {}
for s, r in kand:
    n = str(r['numer'])
    if n not in best or r['rok'] > best[n][1]['rok']:
        best[n] = (s, r)

def premiera(nr, rok):
    p = (sety.get(nr) or {}).get('premiera')
    return str(p) if p else f'{rok}-00'

# od najnowszych: premiera malejąco, potem rocznik, potem cena (droższe wyżej)
data = sorted(best.values(),
              key=lambda x: (premiera(str(x[1]['numer']), x[1]['rok']), x[1]['rok'],
                             float(x[1]['cena_katalogowa'])),
              reverse=True)

ARIAL = 'Arial'
wb = Workbook(); ws = wb.active; ws.title = 'Bez opisu — od najnowszych'
hdr_fill = PatternFill('solid', fgColor='1F3864')
hdr_font = Font(name=ARIAL, size=11, bold=True, color='FFFFFF')
border = Border(bottom=Side(style='thin', color='BFBFBF'))
band = PatternFill('solid', fgColor='F2F2F2')
nowe = PatternFill('solid', fgColor='FFF2CC')

headers = ['Numer', 'Nazwa', 'Seria', 'Premiera', 'Rok', 'Elementy',
           'Cena katalogowa (zł)', 'Jest w sprzedaży', 'Podstrona']
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill = hdr_fill; cell.font = hdr_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

for i, (seria, r) in enumerate(data, start=2):
    nr = str(r['numer'])
    prem = (sety.get(nr) or {}).get('premiera') or ''
    ws.cell(row=i, column=1, value=nr).alignment = Alignment(horizontal='left')
    ws.cell(row=i, column=2, value=r.get('nazwa'))
    ws.cell(row=i, column=3, value=seria)
    ws.cell(row=i, column=4, value=prem).alignment = Alignment(horizontal='center')
    ws.cell(row=i, column=5, value=r['rok']).alignment = Alignment(horizontal='center')
    el = ws.cell(row=i, column=6, value=r.get('elementy'))
    el.alignment = Alignment(horizontal='center')
    pc = ws.cell(row=i, column=7, value=float(r['cena_katalogowa'])); pc.number_format = '#,##0.00'
    ws.cell(row=i, column=8,
            value='tak' if (feed.get(nr) or {}).get('cena') else '—'
            ).alignment = Alignment(horizontal='center')
    link = ws.cell(row=i, column=9, value=f'/zestaw/{nr}/')
    link.hyperlink = f'https://tylkoklocki.pl/zestaw/{nr}/'
    link.font = Font(name=ARIAL, size=10, color='0563C1', underline='single')
    for c in range(1, 10):
        cl = ws.cell(row=i, column=c)
        if c != 9: cl.font = Font(name=ARIAL, size=10)
        cl.border = border
        if r['rok'] == 2026: cl.fill = nowe
        elif i % 2 == 0:     cl.fill = band

for c, w in enumerate([10, 60, 20, 12, 8, 11, 20, 16, 18], 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[1].height = 30
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:I{len(data)+1}'

ws2 = wb.create_sheet('Metodologia')
per_year = Counter(r['rok'] for _, r in data)
per_ser = Counter(s for s, _ in data)
info = [
    ('ZESTAWY BEZ OPISU REDAKCYJNEGO — od najnowszych po najstarsze', ''),
    ('', ''),
    ('Data wygenerowania', datetime.date.today().isoformat()),
    ('Liczba pozycji', len(data)),
    ('', ''),
    ('CO ZNACZY „BEZ OPISU REDAKCYJNEGO"', ''),
    ('Warunek', 'zestaw nie ma ANI karty Piotra, ANI naszego opisu, ANI pary person'),
    ('Odpowiednik w kolejce redakcyjnej', 'status „do opisania"'),
    ('', ''),
    ('KRYTERIA WEJŚCIA (wszystkie naraz)', ''),
    ('1. Rocznik', '2020–2026'),
    ('2. Status w katalogu', 'dostepny — nadal w sprzedaży'),
    ('3. Cena katalogowa', 'znana, w złotych'),
    ('', ''),
    ('SORTOWANIE', 'data premiery malejąco; przy równej dacie decyduje rocznik, '
                   'a potem cena katalogowa'),
    ('Dlaczego tak', 'świeży zestaw bez opisu traci ruch wtedy, kiedy ludzie go '
                     'szukają — dwa lata później ten sam brak kosztuje dużo mniej'),
    ('Podświetlenie na żółto', 'rocznik 2026'),
    ('', ''),
    ('KOLUMNA „JEST W SPRZEDAŻY"', 'czy zestaw ma dziś ofertę w naszym feedzie '
                                   'sklepowym; „—" znaczy, że nie ma go w żadnym '
                                   'ze śledzonych sklepów'),
    ('', ''),
    ('ŹRÓDŁA', ''),
    ('Numer, nazwa, rok, seria, cena, elementy', 'src/data/katalog.json'),
    ('Data premiery, opisy, persony', 'src/data/sety.json'),
    ('Karty katalogowo-sprzedażowe', 'src/data/karty_setow.json'),
    ('Dostępność w sklepach', 'src/data/oferty_feed.json'),
    ('', ''),
    ('ROZKŁAD WG ROCZNIKA', ''),
]
info += [(str(r), per_year[r]) for r in sorted(per_year, reverse=True)]
info += [('', ''), ('NAJWIĘKSZE ZALEGŁOŚCI WG SERII', '')]
info += [(s, n) for s, n in per_ser.most_common(15)]



# --- ARKUSZ 2: szersze ujecie, tam gdzie sa realne pieniadze ---------------
# Populacja kolejki (2020-26, dostepny, z cena) jest juz prawie w calosci
# opisana. Prawdziwa dziura jest gdzie indziej: zestawy STARSZE albo bez
# statusu "dostepny", ktore mimo to maja DZIS oferte w sklepie - czyli ktos
# moze je kupic, a my nie mamy o nich ani zdania.
szersze = [(s, r) for s, r in rows
           if (feed.get(str(r['numer'])) or {}).get('cena')
           and bez_opisu(str(r['numer']))]
best2 = {}
for s, r in szersze:
    n = str(r['numer'])
    if n not in best2 or (r.get('rok') or 0) > (best2[n][1].get('rok') or 0):
        best2[n] = (s, r)
dane2 = sorted(best2.values(),
               key=lambda x: ((x[1].get('rok') or 0),
                              float((feed.get(str(x[1]['numer'])) or {}).get('cena') or 0)),
               reverse=True)

ws3 = wb.create_sheet('Z ofertą, bez tekstu')
h3 = ['Numer', 'Nazwa', 'Seria', 'Rok', 'Elementy', 'Cena katalogowa (zł)',
      'Najniższa dziś (zł)', 'Sklep', 'Status w katalogu']
for c, h in enumerate(h3, 1):
    cell = ws3.cell(row=1, column=c, value=h)
    cell.fill = hdr_fill; cell.font = hdr_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
for i, (seria, r) in enumerate(dane2, start=2):
    nr = str(r['numer']); f = feed.get(nr) or {}
    ws3.cell(row=i, column=1, value=nr).alignment = Alignment(horizontal='left')
    ws3.cell(row=i, column=2, value=r.get('nazwa'))
    ws3.cell(row=i, column=3, value=seria)
    ws3.cell(row=i, column=4, value=r.get('rok')).alignment = Alignment(horizontal='center')
    ws3.cell(row=i, column=5, value=r.get('elementy')).alignment = Alignment(horizontal='center')
    if r.get('cena_katalogowa'):
        cc = ws3.cell(row=i, column=6, value=float(r['cena_katalogowa'])); cc.number_format = '#,##0.00'
    cr = ws3.cell(row=i, column=7, value=float(f.get('cena'))); cr.number_format = '#,##0.00'
    ws3.cell(row=i, column=8, value=f.get('sklep')).alignment = Alignment(horizontal='center')
    ws3.cell(row=i, column=9, value=r.get('status')).alignment = Alignment(horizontal='center')
    for c in range(1, 10):
        cl = ws3.cell(row=i, column=c)
        cl.font = Font(name=ARIAL, size=10); cl.border = border
        if i % 2 == 0: cl.fill = band
for c, w in enumerate([10, 58, 20, 8, 11, 20, 18, 15, 16], 1):
    ws3.column_dimensions[get_column_letter(c)].width = w
ws3.row_dimensions[1].height = 30
ws3.freeze_panes = 'A2'
ws3.auto_filter.ref = f'A1:I{len(dane2)+1}'

info += [('', ''), ('ARKUSZ „Z OFERTĄ, BEZ TEKSTU"', ''),
         ('Liczba pozycji', len(dane2)),
         ('Co to jest', 'każdy zestaw z katalogu, który ma DZIŚ ofertę w naszym '
                        'feedzie sklepowym i nie ma żadnego tekstu — bez ograniczenia '
                        'rocznika i statusu'),
         ('Dlaczego osobno', 'kolejka redakcyjna obejmuje tylko roczniki 2020–2026 ze '
                             'statusem „dostepny" i znaną ceną katalogową; tam zaległość '
                             'jest już prawie zamknięta. Ta lista pokazuje, gdzie została '
                             'naprawdę duża dziura'),
         ('Sortowanie', 'rocznik malejąco, przy równym roczniku najdroższe pierwsze'),
        ]

for i, (a, b) in enumerate(info, start=1):
    ca = ws2.cell(row=i, column=1, value=a)
    cb = ws2.cell(row=i, column=2, value=b)
    ca.font = Font(name=ARIAL, size=11, bold=(i == 1 or (a and not b and a.isupper())))
    cb.font = Font(name=ARIAL, size=11)
    ca.alignment = Alignment(vertical='top')
    cb.alignment = Alignment(wrap_text=True, vertical='top')
ws2.column_dimensions['A'].width = 46
ws2.column_dimensions['B'].width = 70

out = BAZA + 'materialy/zestawy-bez-opisu.xlsx'
wb.save(out)
print('zapisane:', out, '| pozycji:', len(data))
print('rozklad rocznikow:', dict(sorted(per_year.items(), reverse=True)))
print('top serie:', per_ser.most_common(8))
