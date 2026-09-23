#!/usr/bin/env python3
"""Macierz cen: zestawy w wierszach, sklepy w kolumnach, w komórce cena z datą.

Kto uruchamia: człowiek na żądanie (Marek 23.09.2026: „arkusz z numerami
zestawów i sklepami w kolumnach z zaznaczeniem, w których mamy skąd ceny").
Źródła: src/data/sety.json (oferty per sklep, z datą) i src/data/oferty_feed.json
(migawka feedów). Kolor komórki mówi o świeżości względem sita 14 dni z
src/lib/oferty.js: zielony = pokazywana na stronie, żółty = przeterminowana
(w danych jest, na stronie nie), pusto = brak ceny. Arkusz „Źródła" opisuje,
skąd bierze się cena każdego sklepu i jak często. Wynik: materialy/macierz-cen-sklepy.xlsx
"""
import json, datetime, os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

R = os.path.join(os.path.dirname(__file__), '..')
sety = json.load(open(f'{R}/src/data/sety.json', encoding='utf-8'))
feed_all = json.load(open(f'{R}/src/data/oferty_feed.json', encoding='utf-8'))
feed = feed_all.get('sety', {k: v for k, v in feed_all.items() if not k.startswith('_')})
ceny_baza = json.load(open(f'{R}/src/data/ceny_baza.json', encoding='utf-8'))
DZIS = datetime.date.today()
SITO_DNI = 14

SKLEPY = [
    ('lego', 'LEGO.com', 'listing lego.pl przez Firecrawl (lego-ceny.mjs)', 'wtorek (Routine „Dane wt")', 'Dane wt'),
    ('mediaexpert', 'Media Expert', 'feed produktowy (feedy-lego.py)', 'codziennie', 'Łowca'),
    ('planetaklockow', 'Planeta Klocków', 'feed nokaut.xml + kontrola OutOfStock na karcie (feedy-lego.py)', 'codziennie', 'Łowca'),
    ('allegro', 'Allegro', 'feed „tylko LEGO" 39967a61 (feedy-lego.py)', 'codziennie', 'Łowca'),
    ('empik', 'Empik', 'zrzut z przeglądarki (Cowork, skill klocki-ceny-empik) → empik-import.mjs', 'tygodniowo, ręcznie', 'Marek + Cowork'),
    ('smyk', 'Smyk', 'strony produktów (smyk-odswiez.mjs)', 'wtorek i piątek', 'Dane wt / Łowca'),
    ('ceneo', 'Ceneo', 'feed Tradedoubler (ceneo-feed.mjs) – najniższa oferta w porównywarce', 'wtorek', 'Dane wt'),
    ('lidl', 'Lidl', 'feed Tradedoubler (importer „wtorkowy" wołany codziennie)', 'codziennie', 'Łowca'),
    ('xkom', 'x-kom', 'BRAK FEEDU – SalesMasters nie daje feedu, strony blokują ruch serwerowy; ceny tylko ręcznie (mailing partnera, Cowork) z polem wazne_do', 'nieregularnie', 'Marek / sesja Code'),
]
KODY = [k for k, *_ in SKLEPY]

def cena_z_sety(nr, sklep):
    for o in (sety.get(nr) or {}).get('oferty', []) or []:
        if o.get('sklep') == sklep and isinstance(o.get('cena'), (int, float)):
            return o['cena'], o.get('data'), o.get('wazne_do')
    return None

def cena_z_feedu(nr, sklep):
    w = feed.get(nr) or {}
    c = (w.get('oferty') or {}).get(sklep)
    if isinstance(c, (int, float)) and c > 0:
        return c, (w.get('daty') or {}).get(sklep) or w.get('data'), None
    return None

def swieza(data, wazne_do):
    if wazne_do:
        try:
            if DZIS > datetime.date.fromisoformat(wazne_do): return False
        except ValueError: pass
    if not data: return True
    try: return (DZIS - datetime.date.fromisoformat(data[:10])).days <= SITO_DNI
    except ValueError: return True

numery = sorted({n for n in sety if not n.startswith('_')} | {n for n in feed if not n.startswith('_')}, key=lambda x: (len(x), x))

wb = Workbook()
ws = wb.active; ws.title = 'Macierz cen'
ZIEL = PatternFill('solid', fgColor='C6EFCE'); ZOLTY = PatternFill('solid', fgColor='FFEB9C'); SZARY = PatternFill('solid', fgColor='EDEDED')
naglowek = ['Numer', 'Nazwa', 'Seria', 'RRP', 'Ile sklepów (świeże)'] + [n for _, n, *_ in SKLEPY]
ws.append(naglowek)
ws.append(['', '', '', '', 'źródło →'] + [z for _, _, z, *_ in SKLEPY])
ws.append(['', '', '', '', 'częstotliwość →'] + [c for _, _, _, c, _ in SKLEPY])
for cell in ws[1]: cell.font = Font(bold=True)
for r in (2, 3):
    for cell in ws[r]: cell.font = Font(italic=True, size=9); cell.alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[2].height = 60
statystyka = {k: {'swieze': 0, 'stare': 0} for k in KODY}
for nr in numery:
    s = sety.get(nr) or {}
    rrp = s.get('cena_katalogowa') or (ceny_baza.get(nr) or {}).get('cena_katalogowa')
    wiersz = [nr, s.get('nazwa') or (feed.get(nr) or {}).get('nazwa') or '', s.get('seria') or '', rrp, 0]
    kolory = []
    for k in KODY:
        tr = cena_z_sety(nr, k) or cena_z_feedu(nr, k)
        if not tr:
            wiersz.append(''); kolory.append(None); continue
        c, d, wd = tr
        ok = swieza(d, wd)
        wiersz.append(f'{c:.2f}'.replace('.', ',') + (f' ({d[8:10]}.{d[5:7]})' if d else '') + (f' do {wd[8:10]}.{wd[5:7]}' if wd else ''))
        kolory.append(ZIEL if ok else ZOLTY)
        statystyka[k]['swieze' if ok else 'stare'] += 1
        if ok: wiersz[4] += 1
    ws.append(wiersz)
    r = ws.max_row
    for i, kol in enumerate(kolory):
        cell = ws.cell(row=r, column=6 + i)
        cell.fill = kol or SZARY
ws.freeze_panes = 'F4'
for i, w in enumerate([9, 38, 16, 9, 10] + [18] * len(SKLEPY), start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.auto_filter.ref = f'A1:{get_column_letter(len(naglowek))}{ws.max_row}'

zr = wb.create_sheet('Źródła')
zr.append(['Sklep', 'Skąd cena', 'Jak często', 'Kto', 'Zestawów ze świeżą ceną', 'Przeterminowane (w danych, nie na stronie)', 'Afiliacja'])
afil = {'lego': 'Tradedoubler? – patrz afiliacje_rejestr.json', 'mediaexpert': 'Performers', 'planetaklockow': 'webePartners', 'allegro': 'Allegro (bezpośrednio)', 'empik': 'Tradedoubler', 'smyk': 'Tradedoubler', 'ceneo': 'Ceneo Program Partnerski', 'lidl': 'Tradedoubler', 'xkom': 'SalesMasters (kod uniwersalny sm=)'}
for k, n, z, c, kto in SKLEPY:
    zr.append([n, z, c, kto, statystyka[k]['swieze'], statystyka[k]['stare'], afil.get(k, '')])
for cell in zr[1]: cell.font = Font(bold=True)
for i, w in enumerate([16, 70, 24, 18, 12, 14, 30], start=1): zr.column_dimensions[get_column_letter(i)].width = w
zr.append([]); zr.append([f'Stan na {DZIS.isoformat()}. Zielony = cena pokazywana na stronie (sito {SITO_DNI} dni z src/lib/oferty.js), żółty = przeterminowana, pusto = brak. Kolumna „Afiliacja" wg afiliacje_rejestr.json – status sprawdzać tam.'])

out = f'{R}/materialy/macierz-cen-sklepy.xlsx'
wb.save(out)
print(f'zestawów: {len(numery)} | ' + ' | '.join(f"{k}: {v['swieze']} św./{v['stare']} st." for k, v in statystyka.items()))
print('zapisano', os.path.relpath(out, R))
