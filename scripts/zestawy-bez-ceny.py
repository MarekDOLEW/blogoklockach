"""Zestawy w sprzedazy, ktore maja podstrone w serwisie, ale nie znamy ich ceny katalogowej.

Na hubie takiego zestawu stoi wprost "cena katalogowa nieogloszona (podamy, gdy
LEGO ja poda)". Ten listing zbiera je w jednym miejscu z gotowymi linkami, zeby
dalo sie je szybko obejsc i uzupelnic, gdy LEGO ogloszi cennik.

Wycofanych (EOL) bez ceny jest kilkaset i to zaszlosc historyczna - ich ceny
katalogowej czesto nie da sie juz ustalic, wiec nie trafiaja do listingu.

Uruchomienie: python3 scripts/zestawy-bez-ceny.py
"""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

REPO = '/home/user/blogoklockach'
czytaj = lambda p: json.load(open(f'{REPO}/src/data/{p}', encoding='utf-8'))

kat = czytaj('katalog.json')
sety = czytaj('sety.json')
karty = czytaj('karty_setow.json')
rrp = czytaj('rrp_potwierdzone.json')
ceny_baza = czytaj('ceny_baza.json')
feed = czytaj('oferty_feed.json').get('sety', {})
redirects = {k: v for k, v in czytaj('redirects.json').items() if isinstance(v, dict)}
wycofania = {w['numer']: w for w in czytaj('wycofania.json')['wycofania']}

katidx = {}
for seria, lista in kat.items():
    if seria == '_meta' or not isinstance(lista, list):
        continue
    for s in lista:
        katidx[str(s['numer'])] = dict(s, seria=seria)

# ── ktore numery maja podstrone (lustro src/lib/huby.js) ──
def ma_cene_z_feedu(w):
    if not isinstance(w, dict):
        return False
    return bool(w.get('cena') or any(isinstance(c, (int, float)) and c > 0 for c in (w.get('oferty') or {}).values()))

huby = set(sety) | set(wycofania)
for nr in set(feed) | {n for m in redirects.values() for n in m}:
    if nr in huby or (nr not in katidx and nr not in wycofania):
        continue
    if ma_cene_z_feedu(feed.get(nr)) or any(m.get(nr) for m in redirects.values()):
        huby.add(nr)
for nr, s in katidx.items():
    if s.get('status') == 'dostepny' and (s.get('cena_katalogowa') or rrp.get(nr, {}).get('cena')):
        huby.add(nr)
    if karty.get(nr):
        huby.add(nr)


def cena(nr):
    return (rrp.get(nr, {}).get('cena')
            or katidx.get(nr, {}).get('cena_katalogowa')
            or (sety.get(nr) or {}).get('cena_katalogowa')
            or (ceny_baza.get(nr) or {}).get('cena_katalogowa'))


MIESIACE = {'stycznia': '01', 'lutego': '02', 'marca': '03', 'kwietnia': '04', 'maja': '05',
            'czerwca': '06', 'lipca': '07', 'sierpnia': '08', 'września': '09',
            'października': '10', 'listopada': '11', 'grudnia': '12'}


def premiera(nr):
    z_setow = (sety.get(nr) or {}).get('premiera')
    if z_setow:
        return z_setow
    surowa = (karty.get(nr) or {}).get('metryka', {}).get('Premiera') or ''
    czesci = surowa.split()
    for i, slowo in enumerate(czesci):
        if slowo.lower() in MIESIACE and i + 1 < len(czesci) and czesci[i + 1].isdigit():
            return f'{czesci[i + 1]}-{MIESIACE[slowo.lower()]}'
    return ''


def nazwa(nr):
    for zrodlo in ((karty.get(nr) or {}).get('nazwa'), katidx.get(nr, {}).get('nazwa'),
                   (sety.get(nr) or {}).get('nazwa'), (wycofania.get(nr) or {}).get('nazwa')):
        if zrodlo and zrodlo != '{?}':
            return zrodlo
    return '(brak nazwy)'


numery = sorted(n for n in huby if not cena(n) and katidx.get(n, {}).get('status') == 'dostepny')
eol_bez_ceny = sum(1 for n in huby if not cena(n) and katidx.get(n, {}).get('status') != 'dostepny')

ARIAL = 'Arial'
wb = Workbook()
ws = wb.active
ws.title = 'Zestawy bez ceny'

hdr_fill = PatternFill('solid', fgColor='1F3864')
hdr_font = Font(name=ARIAL, size=11, bold=True, color='FFFFFF')
band = PatternFill('solid', fgColor='F2F2F2')
border = Border(bottom=Side(style='thin', color='BFBFBF'))
link_font = Font(name=ARIAL, size=10, color='1155CC', underline='single')

headers = ['Numer', 'Nazwa', 'Seria', 'Rok', 'Elementy', 'Premiera',
           'Karta Piotra', 'Adres podstrony', 'Brickset']
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

for i, nr in enumerate(numery, start=2):
    r = katidx.get(nr, {})
    ws.cell(row=i, column=1, value=nr).alignment = Alignment(horizontal='left')
    ws.cell(row=i, column=2, value=nazwa(nr))
    ws.cell(row=i, column=3, value=r.get('seria', ''))
    ws.cell(row=i, column=4, value=r.get('rok')).alignment = Alignment(horizontal='center')
    ws.cell(row=i, column=5, value=r.get('elementy')).alignment = Alignment(horizontal='center')
    ws.cell(row=i, column=6, value=premiera(nr)).alignment = Alignment(horizontal='center')
    ws.cell(row=i, column=7, value='tak' if nr in karty else '').alignment = Alignment(horizontal='center')

    adres = f'https://tylkoklocki.pl/zestaw/{nr}/'
    kom = ws.cell(row=i, column=8, value=adres)
    kom.hyperlink = adres
    kom_bs = ws.cell(row=i, column=9, value=f'{nr}-1')
    kom_bs.hyperlink = f'https://brickset.com/sets/{nr}-1'

    for kol in range(1, len(headers) + 1):
        cl = ws.cell(row=i, column=kol)
        cl.font = link_font if kol in (8, 9) else Font(name=ARIAL, size=10)
        cl.border = border
        if i % 2 == 0:
            cl.fill = band
        if kol == 2:
            cl.alignment = Alignment(vertical='center', wrap_text=True)

for kol, szer in enumerate([10, 46, 26, 8, 10, 12, 13, 42, 12], start=1):
    ws.column_dimensions[get_column_letter(kol)].width = szer
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(numery) + 1}'

m = wb.create_sheet('Metodologia')
m.column_dimensions['A'].width = 34
m.column_dimensions['B'].width = 110
m.cell(row=1, column=1, value='Zagadnienie').font = hdr_font
m.cell(row=1, column=1).fill = hdr_fill
m.cell(row=1, column=2, value='Wyjasnienie').font = hdr_font
m.cell(row=1, column=2).fill = hdr_fill
wiersze = [
    ('Co to za lista', 'Zestawy w sprzedazy, ktore maja juz podstrone w serwisie, ale nie znamy '
                       'ich ceny katalogowej. Na hubie stoi przy nich wprost: "cena katalogowa '
                       'nieogloszona (podamy, gdy LEGO ja poda)".'),
    ('Skad brak ceny', 'Najczesciej sa to zestawy promocyjne (prezent za zakupy), ekskluzywy '
                       'LEGO House / LEGOLAND, polybagi i swieze zapowiedzi, dla ktorych LEGO '
                       'nie oglosilo jeszcze polskiego cennika.'),
    ('Co zrobic', 'Gdy cena sie pojawi, wpisac ja do src/data/rrp_potwierdzone.json (ma '
                  'pierwszenstwo) albo do katalog.json. Hub, tabele cen i kolejka redakcyjna '
                  'zaktualizuja sie same przy najblizszym buildzie.'),
    ('Adres podstrony', 'Klikalny link do huba zestawu na tylkoklocki.pl.'),
    ('Wycofane bez ceny', f'Poza ta lista jest jeszcze {eol_bez_ceny} zestawow wycofanych bez znanej '
                          'ceny katalogowej. To zaszlosc historyczna - czesto nie da sie juz ustalic '
                          'polskiego RRP - wiec nie traktujemy tego jako zadania do zrobienia.'),
    ('Kiedy odswiezyc', 'Po kazdym imporcie kart i po kazdym uzupelnieniu cen: '
                        'python3 scripts/zestawy-bez-ceny.py'),
]
for i, (a, b) in enumerate(wiersze, start=2):
    m.cell(row=i, column=1, value=a).font = Font(name=ARIAL, size=10, bold=True)
    kb = m.cell(row=i, column=2, value=b)
    kb.font = Font(name=ARIAL, size=10)
    kb.alignment = Alignment(vertical='top', wrap_text=True)
    m.row_dimensions[i].height = 34

sciezka = f'{REPO}/materialy/zestawy-bez-ceny.xlsx'
wb.save(sciezka)
print('zapisano', sciezka, 'pozycji:', len(numery))
print('z karta Piotra:', sum(1 for n in numery if n in karty))
print('wycofanych bez ceny (poza lista):', eol_bez_ceny)
